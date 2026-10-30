# Builds CSV files and one Excel workbook from imports/match-data/*.json (data captured from screenshots the player supplied).
import json,csv,glob,os
from openpyxl import Workbook
out='imports/match-data';wb=Workbook();wb.remove(wb.active);T={}
def add(n,h,rows):T.setdefault(n,[h,[]])[1].extend(rows)
for f in sorted(glob.glob(out+'/*.json')):
    d=json.load(open(f));m=d['meta'];mid=m['match_id']
    add('match',list(m.keys()),[list(m.values())])
    for i,inn in enumerate(d['inn']):
        for k,r in enumerate(inn['bat'] or []):add('batting',['match_id','team','position','batter','runs','balls','fours','sixes','strike_rate','minutes','dismissal'],[[mid,inn['team'],k+1,r[0],r[1],r[2],r[3],r[4],r[5],r[6],r[7]]])
        add('innings',['match_id','team','score','overs','extras','crr','yet_to_bat','batting_card_captured'],[[mid,inn['team'],inn['score'],inn['ov'],inn.get('extras',''),inn.get('crr',''),inn.get('ytb',''),'yes' if inn['bat'] else 'no']])
    ph=d['phases']
    for r in ph['rows']:add('phases',['match_id','phase','metric',ph['A'],ph['B']],[[mid,ph['tabs'],r[0],r[1],r[2]]])
    for o in d['turning']['list']:add('turning_overs',['match_id','batting_team','over','runs','score','balls','batters','bowler'],[[mid,d['turning']['team'],o['ov'],o['runs'],o['sc'],' '.join(o['balls']),o['bat'],o['bowl']]])
    for o in d['overs']['list']:add('overs',['match_id','batting_team','over','score','bowler_to_batters','balls','runs'],[[mid,d['overs']['team'],o[0],o[1],o[2],' '.join(o[3]),o[4]]])
    for k,r in enumerate(d['squads']['rows']):add('squads',['match_id','row',d['squads']['A'],d['squads']['B']],[[mid,k+1,r[0],r[1]]])
    for k,r in enumerate(d['mvp']):add('mvp',['match_id','rank','player','team','mvp_points'],[[mid,k+1,r[0],r[1],r[2]]])
for n,(h,rows) in T.items():
    with open(f'{out}/{n}.csv','w',newline='',encoding='utf-8-sig') as fh:w=csv.writer(fh);w.writerow(h);w.writerows(rows)
    ws=wb.create_sheet(n);ws.append(h);[ws.append(r) for r in rows]
    for c in ws[1]:c.font=c.font.copy(bold=True)
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(40,max(len(str(c.value or '')) for c in col)+2)
wb.save(out+'/scoring-heroes-match-data.xlsx');print({n:len(r[1]) for n,r in T.items()})
